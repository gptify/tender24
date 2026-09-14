import io
import datetime
from typing import Dict, Any, Optional, List
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

try:
    from src.tender_ai.config import DEFAULT_COMPANY_PROFILE
except ImportError:
    from config import DEFAULT_COMPANY_PROFILE

class DocxExporter:
    """Generates formal, production-grade Microsoft Word (.docx) documents 
    compliant with Uzbekistan state, international NGO/UN, and private corporate procurement standards."""

    @staticmethod
    def _set_cell_background(cell, fill_hex: str):
        """Sets background color of a table cell."""
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    @staticmethod
    def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        """Sets inner padding for a table cell."""
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    @classmethod
    def _apply_base_styling(cls, doc: docx.Document):
        """Applies Uzbekistan standard document styling (Times New Roman, 1.15 line spacing)."""
        section = doc.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.7)

        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        font.color.rgb = RGBColor(17, 24, 39)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_after = Pt(6)

    @classmethod
    def create_technical_proposal_docx(
        cls,
        tender_summary: Dict[str, Any],
        proposal_text: str,
        company_profile: Optional[Dict[str, Any]] = None,
        qualifications: Optional[List[Dict[str, Any]]] = None,
        language: str = "uz"
    ) -> bytes:
        """Generates a comprehensive, formally styled Technical Proposal in DOCX format across languages."""
        profile = company_profile or DEFAULT_COMPANY_PROFILE
        doc = docx.Document()
        cls._apply_base_styling(doc)

        # 1. Company Letterhead (Header)
        header_table = doc.add_table(rows=1, cols=2)
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        header_table.autofit = False

        left_cell = header_table.cell(0, 0)
        left_cell.width = Inches(4.0)
        p_left = left_cell.paragraphs[0]
        run_comp = p_left.add_run(f'"{profile.get("name")}"\n')
        run_comp.bold = True
        run_comp.font.size = Pt(13)
        run_comp.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy

        run_desc = p_left.add_run(f"{profile.get('description', '')}\n")
        run_desc.font.size = Pt(9.5)
        run_desc.font.color.rgb = RGBColor(100, 116, 139)

        addr_str = "Address: Tashkent, Uzbekistan | Tel: +998 (90) 000-00-00\nE-mail: info@gptify.uz | Web: https://gptify.uz" if language == "en" else "Manzil: Toshkent sh. | Tel: +998 (90) 000-00-00\nE-mail: info@gptify.uz | Veb: https://gptify.uz"
        run_contact = p_left.add_run(addr_str)
        run_contact.font.size = Pt(8.5)
        run_contact.font.italic = True
        run_contact.font.color.rgb = RGBColor(148, 163, 184)

        right_cell = header_table.cell(0, 1)
        right_cell.width = Inches(2.8)
        p_right = right_cell.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        today_str = datetime.datetime.now().strftime("%d.%m.%Y")
        
        customer_name = tender_summary.get("customer", "Evaluation Committee")
        if language == "en":
            dest_str = f"To: Evaluation Committee of {customer_name}\n\nDate: {today_str}\nRef No: TP-{datetime.datetime.now().strftime('%m%d')}/01"
        elif language == "ru":
            dest_str = f"Кому: Закупочной комиссии {customer_name}\n\nДата: {today_str} г.\nИсх. №: ТП-{datetime.datetime.now().strftime('%m%d')}/01"
        else:
            dest_str = f"Kimgа: {customer_name}\nTanlov va xarid komissiyasiga\n\nSana: {today_str}-yil\nChiqish №: TP-{datetime.datetime.now().strftime('%m%d')}/01"
            
        run_dest = p_right.add_run(dest_str)
        run_dest.font.size = Pt(10)
        run_dest.font.color.rgb = RGBColor(71, 85, 105)

        # Divider line
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_after = Pt(12)
        p_div_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="1E3A8A"/></w:pBdr>')
        p_div._p.get_or_add_pPr().append(p_div_border)

        # 2. Document Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(4)
        
        if language == "en":
            title_text = "TECHNICAL PROPOSAL & RFP RESPONSE\n(INTERNATIONAL PROCUREMENT STANDARD)"
        elif language == "ru":
            title_text = "ТЕХНИЧЕСКОЕ ПРЕДЛОЖЕНИЕ\n(ОФИЦИАЛЬНАЯ ЗАЯВКА НА УЧАСТИЕ)"
        else:
            title_text = "TEXNIK TAKLIF\n(TEXNICHESKOYE PREDLOZHENIYE)"
            
        run_title = p_title.add_run(title_text)
        run_title.bold = True
        run_title.font.size = Pt(15)
        run_title.font.color.rgb = RGBColor(30, 58, 138)

        lot_title = tender_summary.get("title", "Tender")
        lot_num = tender_summary.get("lot_number", "")
        p_lot = doc.add_paragraph()
        p_lot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_lot.paragraph_format.space_after = Pt(16)
        
        lot_lbl = "Lot Ref" if language == "en" else ("Лот №" if language == "ru" else "Lot №")
        subject_lbl = "Subject" if language == "en" else ("Тема" if language == "ru" else "Mavzu")
        run_lot = p_lot.add_run(f'{lot_lbl}: {lot_num}\n{subject_lbl}: "{lot_title}"')
        run_lot.bold = True
        run_lot.font.size = Pt(11)
        run_lot.font.color.rgb = RGBColor(51, 65, 85)

        # 3. Add Sections based on proposal_text
        for line in proposal_text.split("\n"):
            line_str = line.strip()
            if not line_str:
                continue

            if line_str.startswith("# ") or line_str.startswith("## "):
                clean_h = line_str.lstrip("#").strip()
                p_h = doc.add_paragraph()
                p_h.paragraph_format.space_before = Pt(14)
                p_h.paragraph_format.space_after = Pt(4)
                r_h = p_h.add_run(clean_h)
                r_h.bold = True
                r_h.font.size = Pt(13)
                r_h.font.color.rgb = RGBColor(30, 58, 138)
            elif line_str.startswith("### "):
                clean_h = line_str.lstrip("#").strip()
                p_h = doc.add_paragraph()
                p_h.paragraph_format.space_before = Pt(10)
                p_h.paragraph_format.space_after = Pt(2)
                r_h = p_h.add_run(clean_h)
                r_h.bold = True
                r_h.font.size = Pt(11.5)
                r_h.font.color.rgb = RGBColor(15, 23, 42)
            elif line_str.startswith("- ") or line_str.startswith("* "):
                p_b = doc.add_paragraph(style='List Bullet')
                p_b.paragraph_format.space_after = Pt(3)
                cls._add_formatted_text(p_b, line_str[2:])
            elif line_str.startswith("1. ") or line_str.startswith("2. ") or line_str.startswith("3. "):
                p_num = doc.add_paragraph()
                p_num.paragraph_format.space_after = Pt(3)
                p_num.paragraph_format.left_indent = Inches(0.2)
                cls._add_formatted_text(p_num, line_str)
            elif not line_str.startswith("|"):
                p_body = doc.add_paragraph()
                cls._add_formatted_text(p_body, line_str)

        # 4. Detailed Compliance Matrix Table
        if qualifications and len(qualifications) > 0:
            doc.add_paragraph().paragraph_format.space_before = Pt(12)
            p_mat_h = doc.add_paragraph()
            
            matrix_title = "TECHNICAL COMPLIANCE MATRIX (TOR COMPLIANCE)" if language == "en" else ("МАТРИЦА СООТВЕТСТВИЯ ТРЕБОВАНИЯМ ТЗ" if language == "ru" else "TEXNIK TOPSHIRIQ TALABLARI BO'YICHA MUVOFIQLIK JADVALI")
            r_mat = p_mat_h.add_run(matrix_title)
            r_mat.bold = True
            r_mat.font.size = Pt(12)
            r_mat.font.color.rgb = RGBColor(30, 58, 138)

            table = doc.add_table(rows=1, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            if language == "en":
                hdr_titles = ["№ & Category", "TOR Specification / Requirement", "Bidder's Solution & Compliance Status"]
            elif language == "ru":
                hdr_titles = ["№ и Категория", "Требование ТЗ Заказчика", "Предлагаемое решение и статус"]
            else:
                hdr_titles = ["№ va Yo'nalish", "Buyurtmachi Texnik Talabi", "Ishtirokchi Taklifi va Muvofiqligi"]
                
            widths = [Inches(1.8), Inches(2.7), Inches(2.3)]
            for i, title in enumerate(hdr_titles):
                cell = table.rows[0].cells[i]
                cell.width = widths[i]
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(title)
                r.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(255, 255, 255)
                cls._set_cell_background(cell, "1E3A8A")
                cls._set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

            for idx, q in enumerate(qualifications, 1):
                row_cells = table.add_row().cells
                cat = q.get("category", "Talab")
                req = q.get("requirement", "")

                for i in range(3):
                    row_cells[i].width = widths[i]
                    cls._set_cell_margins(row_cells[i], top=80, bottom=80, left=120, right=120)
                    if idx % 2 == 0:
                        cls._set_cell_background(row_cells[i], "F8FAFC")

                p0 = row_cells[0].paragraphs[0]
                p0.add_run(f"{idx}. {cat}").font.size = Pt(9.5)

                p1 = row_cells[1].paragraphs[0]
                p1.add_run(req).font.size = Pt(9.5)

                p2 = row_cells[2].paragraphs[0]
                comp_note = "✅ Fully Compliant. Committed to deliver within specified SLA and timeline." if language == "en" else ("✅ Полностью соответствует. Гарантируется в установленном объеме и сроках." if language == "ru" else "✅ To'liq muvofiq keladi. Talab qilingan hajm va muddatda kafolatlanadi.")
                r_match = p2.add_run(comp_note)
                r_match.font.size = Pt(9.5)
                r_match.font.color.rgb = RGBColor(16, 120, 70)

        # 5. Formal Signature & Stamp Block
        doc.add_paragraph().paragraph_format.space_before = Pt(16)
        sig_table = doc.add_table(rows=1, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_table.autofit = False

        s_left = sig_table.cell(0, 0)
        s_left.width = Inches(4.0)
        p_sig_l = s_left.paragraphs[0]
        
        rep_lbl = f'Authorized Signatory for "{profile.get("name")}":\n\n\n__________________ {profile.get("founder")}\n' if language == "en" else (f'Руководитель "{profile.get("name")}":\n\n\n__________________ {profile.get("founder")}\n' if language == "ru" else f'"{profile.get("name")}" Rahbari:\n\n\n__________________ {profile.get("founder")}\n')
        p_sig_l.add_run(rep_lbl).font.size = Pt(11)
        r_sub = p_sig_l.add_run("(signature, date)" if language == "en" else ("(подпись, дата)" if language == "ru" else "(imzo, sana)"))
        r_sub.font.size = Pt(8.5)
        r_sub.font.color.rgb = RGBColor(100, 116, 139)

        s_right = sig_table.cell(0, 1)
        s_right.width = Inches(2.8)
        p_sig_r = s_right.paragraphs[0]
        p_sig_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_sig_r.add_run("\n\nL.S. (Corporate Stamp / M.P.)\n\n" if language == "en" else "\n\nM.O'. (Muhr o'rni)\n\n").font.size = Pt(11)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    def create_guarantee_letter_docx(
        cls,
        tender_summary: Dict[str, Any],
        guarantee_text: str,
        company_profile: Optional[Dict[str, Any]] = None,
        language: str = "uz"
    ) -> bytes:
        """Generates formal Guarantee Letter / Bid Submission Letter in DOCX format."""
        profile = company_profile or DEFAULT_COMPANY_PROFILE
        doc = docx.Document()
        cls._apply_base_styling(doc)

        p_comp = doc.add_paragraph()
        p_comp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_comp = p_comp.add_run(f'"{profile.get("name").upper()}"\n')
        r_comp.bold = True
        r_comp.font.size = Pt(14)
        r_comp.font.color.rgb = RGBColor(30, 58, 138)

        r_sub = p_comp.add_run(f"{profile.get('description')}\nTashkent, Uzbekistan | info@gptify.uz | Tel: +998 (90) 000-00-00\n")
        r_sub.font.size = Pt(9.5)
        r_sub.font.color.rgb = RGBColor(100, 116, 139)

        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_after = Pt(10)
        p_div_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="1E3A8A"/></w:pBdr>')
        p_div._p.get_or_add_pPr().append(p_div_border)

        customer_name = tender_summary.get("customer", "Evaluation Committee")
        today_str = datetime.datetime.now().strftime("%d.%m.%Y")
        
        p_to = doc.add_paragraph()
        p_to.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_to.paragraph_format.space_after = Pt(14)
        
        if language == "en":
            r_to = p_to.add_run(f"To: Evaluation Committee of {customer_name}\n\nDate: {today_str}\nRef No: GL-{datetime.datetime.now().strftime('%m%d')}/02")
            g_title = "OFFICIAL BID SUBMISSION & GUARANTEE LETTER"
        elif language == "ru":
            r_to = p_to.add_run(f"Кому: Закупочной комиссии {customer_name}\n\nДата: {today_str} г.\nИсх. №: ГП-{datetime.datetime.now().strftime('%m%d')}/02")
            g_title = "ГАРАНТИЙНОЕ ПИСЬМО"
        else:
            r_to = p_to.add_run(f"Kimgа: {customer_name}\nTanlov va xarid komissiyasiga\n\nSana: {today_str}-yil\nChiqish №: KX-{datetime.datetime.now().strftime('%m%d')}/02")
            g_title = "KAFOLAT XATI"
            
        r_to.font.size = Pt(10.5)

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(14)
        r_title = p_title.add_run(g_title)
        r_title.bold = True
        r_title.font.size = Pt(16)
        r_title.font.color.rgb = RGBColor(30, 58, 138)

        lines = guarantee_text.split("\n")
        for line in lines:
            line_str = line.strip()
            if not line_str or line_str.startswith("#") or "KAFOLAT XATI" in line_str or "GUARANTEE" in line_str:
                continue
            if line_str.startswith("1. ") or line_str.startswith("2. ") or line_str.startswith("3. ") or line_str.startswith("4. "):
                p_item = doc.add_paragraph()
                p_item.paragraph_format.space_after = Pt(4)
                p_item.paragraph_format.left_indent = Inches(0.2)
                cls._add_formatted_text(p_item, line_str)
            else:
                p_body = doc.add_paragraph()
                cls._add_formatted_text(p_body, line_str)

        # Signature Block
        doc.add_paragraph().paragraph_format.space_before = Pt(20)
        s_table = doc.add_table(rows=1, cols=2)
        s_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        s_table.autofit = False

        c_l = s_table.cell(0, 0)
        c_l.width = Inches(4.0)
        rep_lbl = f'Authorized Signatory for "{profile.get("name")}":\n\n\n__________________ {profile.get("founder")}' if language == "en" else f'"{profile.get("name")}" Rahbari:\n\n\n__________________ {profile.get("founder")}'
        c_l.paragraphs[0].add_run(rep_lbl).font.size = Pt(11)

        c_r = s_table.cell(0, 1)
        c_r.width = Inches(2.8)
        c_r.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        c_r.paragraphs[0].add_run("\n\nL.S. (Stamp / M.O'.)\n\n" if language == "en" else "\n\nM.O'. (Muhr o'rni)\n\n").font.size = Pt(11)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    def create_checklist_docx(
        cls,
        tender_summary: Dict[str, Any],
        company_profile: Optional[Dict[str, Any]] = None,
        language: str = "uz"
    ) -> bytes:
        """Generates 10-point Submission Checklist in DOCX format."""
        profile = company_profile or DEFAULT_COMPANY_PROFILE
        doc = docx.Document()
        cls._apply_base_styling(doc)

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(4)
        
        c_title = "MANDATORY BID SUBMISSION CHECKLIST" if language == "en" else ("КОНТРОЛЬНЫЙ СПИСОК ДОКУМЕНТОВ" if language == "ru" else "TENDER HUJJATLARI NAZORAT RO'YXATI (CHECKLIST)")
        r_title = p_title.add_run(c_title)
        r_title.bold = True
        r_title.font.size = Pt(15)
        r_title.font.color.rgb = RGBColor(30, 58, 138)

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(14)
        p_sub.add_run(f"Lot: {tender_summary.get('lot_number', 'Noma\'lum')} | {tender_summary.get('customer', '')}\nIshtirokchi: {profile.get('name')}").font.size = Pt(10.5)

        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        if language == "en":
            hdr_titles = ["№", "Document Name", "Source / Legal Form", "Status"]
            docs_list = [
                ("Certificate of Legal Registration", "State One-Stop Shop / QR-code PDF", "[   ] Ready"),
                ("Tax Clearance Certificate", "my.soliq.uz official statement", "[   ] Ready"),
                ("Audited Balance Sheet & Financial Reports", "Past 1-2 fiscal years tax approved", "[   ] Ready"),
                ("Past Contracts & Completion Certificates", "Minimum 2 relevant signed contracts", "[   ] Ready"),
                ("Staffing & Key Personnel Credentials", "Diplomas, resumes and employment records", "[   ] Ready"),
                ("Turnkey Technical Proposal", "Generated by TenderPro AI, signed and stamped", "[   ] Ready"),
                ("Commercial Price Schedule", "Competitive price breakdown within budget", "[   ] Ready"),
                ("Official Letter of Guarantee", "On company letterhead with signature and seal", "[   ] Ready"),
                ("Bid Security / Deposit (if applicable)", "3% bank guarantee or portal deposit receipt", "[   ] Ready"),
                ("Digital Signature / Power of Attorney", "All documents sealed with authorized signature", "[   ] Ready")
            ]
        else:
            hdr_titles = ["№", "Hujjat Nomi", "Manba / Rasmiylashtirish shakli", "Nazorat Belgisi"]
            docs_list = [
                ("Davlat ro'yxatidan o'tganlik Guvohnomasi", "my.gov.uz orqali QR-kodli elektron nusxa", "[   ] Mavjud"),
                ("Soliqdan qarz yo'qligi to'g'risida ma'lumotnoma", "my.soliq.uz shaxsiy kabinetidan yangi olingan", "[   ] Mavjud"),
                ("Moliyaviy hisobot (1-shakl Balans, 2-shakl)", "Oxirgi hisobot davri uchun soliq tasdiqlagan", "[   ] Mavjud"),
                ("O'xshash yo'nalishdagi shartnomalar va aktlar", "Kamida 2 ta to'liq bajarilgan loyiha dalolatnomasi", "[   ] Mavjud"),
                ("Xodimlar malakasi va shtat jadvali", "IT muhandislari diplomi, sertifikatlari va buyruqlari", "[   ] Mavjud"),
                ("Rasmiy Texnik Taklif", "Ushbu tizimda yaratilgan va rahbar imzolagan nusxa", "[   ] Mavjud"),
                ("Tijorat taklifi (Narxlar smetasi)", "Boshlang'ich narxdan 3-7% tushirilgan narx jadvali", "[   ] Mavjud"),
                ("Kafolat xati (Garantiynoe pismo)", "Korxona rasmiy blankida muhr va imzo bilan", "[   ] Mavjud"),
                ("Zaklad to'lovi yoki Bank kafolati", "etender.uzex.uz shaxsiy hisobiga 3% to'lov kvitansiyasi", "[   ] Mavjud"),
                ("Elektron Raqamli Imzo (E-Imzo / ERI)", "Barcha yuklangan PDF fayllarni E-imzo bilan tasdiqlash", "[   ] Mavjud")
            ]

        widths = [Inches(0.6), Inches(2.5), Inches(2.6), Inches(1.1)]
        for i, title in enumerate(hdr_titles):
            cell = table.rows[0].cells[i]
            cell.width = widths[i]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(title)
            r.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(255, 255, 255)
            cls._set_cell_background(cell, "1E3A8A")
            cls._set_cell_margins(cell, top=100, bottom=100, left=100, right=100)

        for idx, (name, source, status) in enumerate(docs_list, 1):
            row = table.add_row().cells
            for i in range(4):
                row[i].width = widths[i]
                cls._set_cell_margins(row[i], top=70, bottom=70, left=90, right=90)
                if idx % 2 == 0:
                    cls._set_cell_background(row[i], "F8FAFC")

            row[0].paragraphs[0].add_run(str(idx)).font.size = Pt(9.5)
            row[1].paragraphs[0].add_run(name).font.size = Pt(9.5)
            row[2].paragraphs[0].add_run(source).font.size = Pt(9.5)
            r_st = row[3].paragraphs[0].add_run(status)
            r_st.font.size = Pt(9.5)
            r_st.bold = True
            r_st.font.color.rgb = RGBColor(30, 58, 138)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def _add_formatted_text(paragraph, text: str):
        """Parses simple markdown bold (**text**) and renders runs."""
        import re
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = paragraph.add_run(part[2:-2])
                r.bold = True
            else:
                paragraph.add_run(part)
